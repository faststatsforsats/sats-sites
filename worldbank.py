"""World Bank Commodity Price Data (the "Pink Sheet"): monthly gold price in USD per troy ounce since 1960, no key.

Page: https://www.worldbank.org/en/research/commodity-markets (updated on the first days of each month).
The monthly history is an Excel file whose link on that page changes once a year, so the build reads the page,
finds the current link, downloads the workbook, and takes the "Gold" column of the "Monthly Prices" sheet.
License: CC BY 4.0 (https://datacatalog.worldbank.org/public-licenses#cc-by).

FRED used to carry the daily LBMA gold fix, but ICE Benchmark Administration data was removed from FRED in 2022,
which is why the monthly Pink Sheet is the source here.
"""

from __future__ import annotations

import io
import re

from .http import SourceError, get_bytes, get_text

SOURCE = {"name": "World Bank Commodity Price Data (The Pink Sheet)", "url": "https://www.worldbank.org/en/research/commodity-markets"}
PAGE = "https://www.worldbank.org/en/research/commodity-markets"
LAST_KNOWN = "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx"


def _workbook_url() -> str:
    try:
        html = get_text(PAGE, fixture="worldbank_page.html")
    except SourceError:
        return LAST_KNOWN
    match = re.search(r'https://thedocs\.worldbank\.org/[^"\']+CMO-Historical-Data-Monthly\.xlsx', html)
    return match.group(0) if match else LAST_KNOWN


def gold_monthly() -> list[tuple[str, float]]:
    """Return [(YYYY-MM-01, usd per troy ounce), ...] oldest first."""
    import openpyxl  # imported here so the site build never needs it

    raw = get_bytes(_workbook_url(), fixture="worldbank_cmo.xlsx")
    try:
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as err:  # noqa: BLE001
        raise SourceError(f"World Bank workbook could not be opened: {err}") from err
    sheet = None
    for name in book.sheetnames:
        if name.strip().lower() == "monthly prices":
            sheet = book[name]
            break
    if sheet is None:
        raise SourceError("World Bank workbook: no 'Monthly Prices' sheet")

    gold_col = None
    points = []
    for row in sheet.iter_rows(values_only=True):
        if gold_col is None:
            for idx, cell in enumerate(row):
                if isinstance(cell, str) and cell.strip().lower() == "gold":
                    gold_col = idx
                    break
            continue
        label = row[0]
        if not isinstance(label, str):
            continue
        match = re.match(r"^(\d{4})M(\d{2})$", label.strip())
        if not match or gold_col >= len(row):
            continue
        value = row[gold_col]
        if isinstance(value, (int, float)):
            points.append((f"{match.group(1)}-{match.group(2)}-01", float(value)))
    if not points:
        raise SourceError("World Bank workbook: no gold rows found")
    return sorted(points)
